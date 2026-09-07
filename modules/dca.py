# ============================================================
#  QuantEdge — Module DCA Long Terme
#  Lit directement simu_invest.xlsm (feuille "data")
#  Python 3.11–3.12
# ============================================================

from typing import Optional, Dict, List, Tuple
import pandas as pd
import numpy as np
import plotly.graph_objects as go
import plotly.express as px
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
import os
import time
import logging
import zipfile
import math
import tempfile
from zoneinfo import ZoneInfo
from datetime import datetime

from config import DCA_PATH, DCA_XLSX_PATH, DCA_DEMO_PATH

logger = logging.getLogger(__name__)

# ── Chemins ──────────────────────────────────────────────────
EXCEL_PATH = DCA_PATH
XLSX_EXCEL_PATH = DCA_XLSX_PATH
DEMO_EXCEL_PATH = DCA_DEMO_PATH

MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_WORKBOOK_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DCA_ROWS = 10_000
MAX_DCA_COLUMNS = 128


class DCADataError(RuntimeError):
    """Raised when a DCA workbook exists but cannot be parsed safely."""


def resolve_dca_source() -> Tuple[Optional[str], str]:
    """Return (path, mode), preferring a validated local workbook if present."""
    for path in (EXCEL_PATH, XLSX_EXCEL_PATH):
        if os.path.exists(path):
            return path, "personal"
    if os.path.exists(DEMO_EXCEL_PATH):
        return DEMO_EXCEL_PATH, "demo"
    return None, "missing"


def _finite_number(value, *, allow_zero: bool = True) -> Optional[float]:
    if value is None or value == "":
        return 0.0 if allow_zero else None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    if number < 0 or (not allow_zero and number == 0):
        return None
    return number

# ── Mapping produits → colonnes data ─────────────────────────
# Colonne paire = valeur totale actuelle
PRODUITS_MAP = {
    "iShares Global Infrastructure Dist":    26,
    "iShares Agribusiness Acc":              28,
    "iShares MSCI India Acc":                30,
    "L&G Artificial Intelligence Acc":       32,
    "iShares MSCI World Small Cap Acc":      34,
    "iShares Core MSCI Europe Acc":          36,
    "iShares Physical Gold Acc":             38,
    "iShares Digital Security Acc":          40,
    "iShares Automation & Robotics Acc":     42,
    "iShares Global Clean Energy Dist":      44,
    "iShares Global Clean Energy Acc":       46,
    "Nvidia":                                48,
    "Xtrackers MSCI World ex USA Acc":       50,
    "iShares Global Water Acc":              52,
    "Global X Uranium Acc":                  54,
    "Cameco":                                56,
    "IonQ":                                  58,
    "iShares S&P 500 Health Care Acc":       60,
    "Vanguard FTSE All World Acc":           62,
    "iShares Global Aggregate Bond EUR Acc": 64,
    "VanEck Global Real Estate Acc":         66,
    "Schneider Electric":                    68,
}

ETF_PRINCIPAUX = {
    "iShares S&P 500 Acc":              {"inv": 2, "val": 8},
    "iShares NASDAQ 100 Acc":           {"inv": 3, "val": 9},
    "HSBC MSCI Emerging Markets Acc":   {"inv": 4, "val": 10},
    "VanEck Defense Acc":               {"inv": 5, "val": 11},
}

# Explicit dashboard aliases. Never infer invested capital from market value.
DASHBOARD_PRODUCT_ALIASES = {
    "iShares Physical Gold Acc (environ)": "iShares Physical Gold Acc",
}

COLOR_BG   = "rgba(0,0,0,0)"
COLOR_GRID = "rgba(128,128,128,0.15)"
COLOR_GREEN = "#22c55e"
COLOR_RED   = "#ef4444"
COLOR_AMBER = "#f59e0b"
COLOR_BLUE  = "#0ea5e9"


# ============================================================
#  LECTURE DU FICHIER EXCEL
# ============================================================


def _validate_workbook_archive(path: str) -> None:
    """Reject oversized or suspiciously expanded XLSX/XLSM archives before openpyxl parses them."""
    try:
        compressed_size = os.path.getsize(path)
    except OSError as exc:
        raise DCADataError(f"Unable to inspect workbook: {exc}") from exc
    if compressed_size > MAX_WORKBOOK_BYTES:
        raise DCADataError(f"Workbook exceeds {MAX_WORKBOOK_BYTES // (1024*1024)} MB upload limit")
    try:
        with zipfile.ZipFile(path) as zf:
            expanded = sum(max(0, int(info.file_size)) for info in zf.infolist())
            if expanded > MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                raise DCADataError(
                    f"Workbook expands beyond {MAX_WORKBOOK_UNCOMPRESSED_BYTES // (1024*1024)} MB safety limit"
                )
    except zipfile.BadZipFile as exc:
        raise DCADataError("Workbook is not a valid XLSX/XLSM archive") from exc


def load_excel(path: str = EXCEL_PATH) -> Optional[pd.DataFrame]:
    """Parse and validate the DCA ``data`` sheet without accepting NaN/inf/negative values.

    Rows are first parsed independently, then sorted chronologically before contribution
    deltas are inferred. This prevents an out-of-order worksheet from silently erasing a
    dated cash flow.
    """
    if not os.path.exists(path):
        return None

    _validate_workbook_archive(path)
    try:
        wb = openpyxl.load_workbook(path, keep_vba=path.lower().endswith(".xlsm"), data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise DCADataError(f"Unable to open DCA workbook: {exc}") from exc

    try:
        if "data" not in wb.sheetnames:
            raise DCADataError("DCA workbook must contain a 'data' sheet")
        ws = wb["data"]

        if ws.max_row > MAX_DCA_ROWS or ws.max_column > MAX_DCA_COLUMNS:
            raise DCADataError(
                f"DCA sheet dimensions exceed safety limits: "
                f"{ws.max_row} rows x {ws.max_column} columns"
            )

        parsed_rows = []
        today = datetime.now(ZoneInfo("Europe/Paris")).date()

        for excel_row, row in enumerate(ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), values_only=True), start=2):
            if not row or row[0] is None:
                continue
            if len(row) <= 11:
                raise DCADataError("DCA 'data' sheet does not contain the expected columns")

            date_value = pd.to_datetime(row[0], errors="coerce")
            if pd.isna(date_value):
                raise DCADataError(f"Invalid date in DCA row {excel_row}")
            date_value = pd.Timestamp(date_value).tz_localize(None)

            # The personal workbook contains dated future template rows. A row with
            # no economic values at all is inert metadata and can be ignored before
            # requiring total_inv.
            if date_value.date() > today and all(v in (None, "") for v in row[1:12]):
                continue

            names = ["sp500_inv", "nasdaq_inv", "em_inv", "defense_inv", "total_inv", "autre_inv",
                     "sp500_val", "nasdaq_val", "em_val", "defense_val", "total_val"]
            parsed = {}
            for offset, name in enumerate(names, start=1):
                raw_value = row[offset]

                if name == "total_inv" and raw_value is None:
                    raise DCADataError(f"Missing {name} in DCA row {excel_row}")

                if (
                    name in {
                        "sp500_val",
                        "nasdaq_val",
                        "em_val",
                        "defense_val",
                    }
                    and raw_value in (None, "")
                ):
                    parsed[name] = np.nan
                    continue

                value = _finite_number(raw_value)

                if value is None:
                    raise DCADataError(
                        f"Invalid/non-finite/negative {name} "
                        f"in DCA row {excel_row}"
                    )

                parsed[name] = value

            entry = {"date": date_value, "excel_row": excel_row, **parsed, "valuation_available": parsed["total_val"] > 0}
            if parsed["total_val"] <= 0:
                entry["total_val"] = np.nan
                for nom in PRODUITS_MAP:
                    entry[nom] = np.nan
            else:
                for nom, col in PRODUITS_MAP.items():
                    idx = col - 1
                    raw = row[idx] if idx < len(row) else None

                    if raw in (None, ""):
                        entry[nom] = np.nan
                        continue

                    value = _finite_number(raw)

                    if value is None:
                        raise DCADataError(
                            f"Invalid/non-finite/negative value for {nom} "
                            f"in DCA row {excel_row}"
                        )

                    entry[nom] = value
            parsed_rows.append(entry)

        if not parsed_rows:
            raise DCADataError("No valid DCA observations found in the 'data' sheet")
    finally:
        wb.close()

    df = pd.DataFrame(parsed_rows).sort_values("date").reset_index(drop=True)
    if df["date"].duplicated().any():
        raise DCADataError("Duplicate DCA snapshot dates are not allowed")

    valuation_available = pd.to_numeric(df["total_val"], errors="coerce").notna()
    future_mask = df["date"].dt.date > today
    historical = df.loc[~future_mask].copy()
    future = df.loc[future_mask].copy()
    if historical.empty:
        if not future.empty:
            future_invested = pd.to_numeric(future["total_inv"], errors="coerce").fillna(0)
            future_valuation = pd.to_numeric(future["total_val"], errors="coerce").notna()
            bad_future = future_valuation | (future_invested > 1e-9)
            if bad_future.any():
                row_num = int(future.loc[bad_future, "excel_row"].iloc[0])
                raise DCADataError(f"Future DCA observation/contribution in row {row_num}")
        raise DCADataError("No valid historical DCA observations found in the 'data' sheet")

    historical_invested = pd.to_numeric(historical["total_inv"], errors="coerce")
    historical_changes = historical_invested.diff()
    if (historical_changes.dropna() < -1e-6).any():
        raise DCADataError("Cumulative total_inv decreases between historical snapshots; XIRR cash flows would be ambiguous")
    historical_contribution = historical_changes.fillna(historical_invested).gt(1e-9)
    historical_valuation = pd.to_numeric(historical["total_val"], errors="coerce").notna()

    # Future workbook rows are never historical observations. A future valuation
    # or a planned capital increase above the latest observed invested amount is
    # rejected. Legacy/template rows with zero valuation and a non-increasing
    # placeholder total_inv are ignored.
    last_historical_invested = float(historical_invested.iloc[-1])
    if not future.empty:
        future_invested = pd.to_numeric(future["total_inv"], errors="coerce")
        future_valuation = pd.to_numeric(future["total_val"], errors="coerce").notna()
        bad_future = future_valuation | (future_invested > last_historical_invested + 1e-9)
        if bad_future.any():
            row_num = int(future.loc[bad_future, "excel_row"].iloc[0])
            raise DCADataError(f"Future DCA observation/contribution in row {row_num}")

    # Preserve past zero-valuation rows only when they carry a dated capital
    # increase. Drop past inert/template rows.
    keep_historical = historical_valuation | historical_contribution
    df = historical.loc[keep_historical].copy().reset_index(drop=True)
    if df.empty:
        raise DCADataError("No valid historical DCA observations found in the 'data' sheet")

    if pd.isna(pd.to_numeric(df["total_val"], errors="coerce").iloc[-1]):
        raise DCADataError("Latest DCA contribution has no subsequent portfolio valuation")

    df = df.drop(columns=["excel_row"], errors="ignore")
    return df


def validate_workbook(path: str) -> pd.DataFrame:
    """Validate workbook structure/data before it is allowed to replace local personal data."""
    df = load_excel(path)
    if df is None or df.empty:
        raise DCADataError("Workbook contains no usable DCA data")
    return df


def save_uploaded_workbook(content: bytes, filename: str) -> str:
    """Validate an upload in a temp file, then atomically replace the matching personal workbook."""
    suffix = __import__("pathlib").Path(filename or "").suffix.lower()
    if len(content) > MAX_WORKBOOK_BYTES:
        raise DCADataError(f"Workbook exceeds {MAX_WORKBOOK_BYTES // (1024*1024)} MB upload limit")
    if suffix not in {".xlsm", ".xlsx"}:
        raise DCADataError("Only .xlsm and .xlsx files are supported")
    target = EXCEL_PATH if suffix == ".xlsm" else XLSX_EXCEL_PATH
    other = XLSX_EXCEL_PATH if suffix == ".xlsm" else EXCEL_PATH
    os.makedirs(os.path.dirname(target), exist_ok=True)
    fd, temp_path = tempfile.mkstemp(prefix="quantedge_dca_", suffix=suffix, dir=os.path.dirname(target))
    try:
        with os.fdopen(fd, "wb") as f:
            f.write(content)
            f.flush()
            os.fsync(f.fileno())
        validate_workbook(temp_path)
        os.replace(temp_path, target)
        if os.path.exists(other):
            os.remove(other)
        return target
    except Exception:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        raise


def get_file_info(path: str = EXCEL_PATH) -> Dict:
    """Infos sur le fichier Excel."""
    if not os.path.exists(path):
        return {"exists": False}
    mtime = os.path.getmtime(path)
    return {
        "exists":   True,
        "modified": pd.Timestamp(mtime, unit='s').strftime("%d/%m/%Y %H:%M"),
        "size_kb":  round(os.path.getsize(path) / 1024, 1),
    }




def calculate_xirr(df: pd.DataFrame) -> Optional[float]:
    """
    Money-weighted annual return for a DCA portfolio.

    Cumulative ``total_inv`` is converted into dated external cash flows:
    contributions are negative investor cash flows and the final portfolio
    value is a positive terminal cash flow. A bisection solver is used so no
    SciPy dependency is required.
    """
    if df is None or df.empty or len(df) < 2:
        return None

    ordered = df.sort_values("date").copy()
    ordered["date"] = pd.to_datetime(ordered["date"])
    invested = pd.to_numeric(ordered["total_inv"], errors="coerce").replace([np.inf, -np.inf], np.nan)
    terminal_values = pd.to_numeric(ordered["total_val"], errors="coerce").replace([np.inf, -np.inf], np.nan)
    if invested.isna().any() or (invested < 0).any():
        return None
    # Intermediate valuation gaps are allowed because contribution timing can still
    # be known. The terminal observation must have a real positive valuation.
    if terminal_values.dropna().empty or pd.isna(terminal_values.iloc[-1]) or (terminal_values.dropna() < 0).any():
        return None

    changes = invested.diff()
    if (changes.dropna() < -1e-9).any():
        return None
    changes.iloc[0] = invested.iloc[0]
    cashflows = []
    for dt, amount in zip(ordered["date"], changes):
        if abs(float(amount)) > 1e-9:
            cashflows.append((dt.to_pydatetime(), -float(amount)))

    final_value = float(terminal_values.iloc[-1])
    if final_value <= 0:
        return None
    cashflows.append((ordered.iloc[-1]["date"].to_pydatetime(), final_value))

    if not any(v < 0 for _, v in cashflows) or not any(v > 0 for _, v in cashflows):
        return None

    t0 = min(dt for dt, _ in cashflows)

    def npv(rate: float) -> float:
        if rate <= -1:
            return float("inf")
        total = 0.0
        for dt, value in cashflows:
            years = (dt - t0).days / 365.2425
            total += value / ((1.0 + rate) ** years)
        return total

    low, high = -0.9999, 1.0
    f_low, f_high = npv(low), npv(high)
    # Expand the positive bound for unusually high returns.
    while f_low * f_high > 0 and high < 1_000:
        high *= 2
        f_high = npv(high)
    if f_low * f_high > 0:
        return None

    for _ in range(200):
        mid = (low + high) / 2
        f_mid = npv(mid)
        if abs(f_mid) < 1e-8:
            return mid
        if f_low * f_mid <= 0:
            high, f_high = mid, f_mid
        else:
            low, f_low = mid, f_mid
    return (low + high) / 2


# ============================================================
#  STATISTIQUES GLOBALES
# ============================================================

def compute_global_stats(df: pd.DataFrame) -> Dict:
    """Compute portfolio statistics with explicit snapshot/contribution semantics."""
    if df is None or df.empty:
        return {}
    ordered = df.sort_values("date").copy()
    last, first = ordered.iloc[-1], ordered.iloc[0]
    total_inv = float(last["total_inv"])
    total_val = float(last["total_val"])
    pnl = total_val - total_inv
    pnl_pct = pnl / total_inv if total_inv > 0 else 0.0

    valued = ordered[pd.to_numeric(ordered["total_val"], errors="coerce").notna()].copy()
    n_observations = len(valued)
    elapsed_days = max(1, int((last["date"] - first["date"]).days))
    duree_mois = max(1.0, elapsed_days / 30.4375)
    changes = pd.to_numeric(ordered["total_inv"], errors="coerce").diff()
    initial_invested = float(ordered.iloc[0]["total_inv"])
    subsequent_contributions = changes.iloc[1:]
    subsequent_contributions = subsequent_contributions[subsequent_contributions > 1e-9]
    contribution_events = int(len(subsequent_contributions))
    gross_contributions = float(subsequent_contributions.sum())
    # Projection prefill excludes the initial funded capital; only subsequent
    # contributions are spread across elapsed time.
    monthly_invested_capital_rate = gross_contributions / duree_mois if duree_mois > 0 else 0.0

    ratio = total_val / total_inv if total_inv > 0 else 1.0
    xirr = calculate_xirr(ordered)

    periods = valued[["date", "total_val", "total_inv"]].copy()
    periods["net_gain"] = periods["total_val"].diff() - periods["total_inv"].diff()
    valid_gains = periods["net_gain"].dropna()
    best_period = round(float(valid_gains.max()), 2) if not valid_gains.empty else None
    worst_period = round(float(valid_gains.min()), 2) if not valid_gains.empty else None

    return {
        "total_inv": round(total_inv, 2), "total_val": round(total_val, 2),
        "pnl": round(pnl, 2), "pnl_pct": round(pnl_pct, 4),
        "n_observations": n_observations, "duree_mois": round(duree_mois, 1),
        "contribution_events": contribution_events,
        "initial_invested": round(initial_invested, 2),
        "gross_contributions": round(gross_contributions, 2),
        "avg_monthly_contribution": round(monthly_invested_capital_rate, 2),
        "avg_monthly_contribution_is_estimate": True,
        "ratio_val_inv": round(ratio, 3),
        "best_period_net_gain": best_period, "worst_period_net_gain": worst_period,
        "date_debut": first["date"].strftime("%d/%m/%Y"),
        "date_fin": last["date"].strftime("%d/%m/%Y"),
        "xirr": round(xirr, 4) if xirr is not None else None,
    }


# ============================================================
#  GRAPHIQUES
# ============================================================

def build_evolution_curve(df: pd.DataFrame) -> go.Figure:
    """Courbe valeur totale vs montant investi dans le temps."""
    plot_df = df[pd.to_numeric(df["total_val"], errors="coerce").notna()].copy()
    fig = go.Figure()

    # Plot invested capital first, then fill the current-value trace to that
    # previous trace. With Plotly, ``fill=tonexty`` on the first trace fills
    # toward zero rather than between invested capital and portfolio value.
    fig.add_trace(go.Scatter(
        x=plot_df["date"], y=plot_df["total_inv"],
        name="Total invested",
        line=dict(color=COLOR_BLUE, width=2, dash="dash"),
        hovertemplate="<b>%{x|%b %Y}</b><br>Investi : %{y:,.0f}€<extra></extra>",
    ))

    fig.add_trace(go.Scatter(
        x=plot_df["date"], y=plot_df["total_val"],
        name="Current value",
        line=dict(color=COLOR_GREEN, width=2.5),
        fill="tonexty",
        fillcolor="rgba(34,197,94,0.15)",
        hovertemplate="<b>%{x|%b %Y}</b><br>Valeur : %{y:,.0f}€<extra></extra>",
    ))

    fig.update_layout(
        title=dict(text="Évolution du portfolio DCA", font=dict(size=15), x=0.02),
        xaxis=dict(showgrid=True, gridcolor=COLOR_GRID),
        yaxis=dict(title="Montant (€)", showgrid=True, gridcolor=COLOR_GRID, ticksuffix="€"),
        plot_bgcolor=COLOR_BG, paper_bgcolor=COLOR_BG,
        margin=dict(l=60, r=20, t=50, b=40),
        hovermode="x unified", height=350,
        legend=dict(orientation="h", x=0, y=1.1),
    )
    return fig


def build_etf_evolution(df: pd.DataFrame, produit: str) -> go.Figure:
    """Évolution d'un ETF spécifique dans le temps."""
    if produit not in df.columns:
        return go.Figure()

    series = df[["date", produit]].copy()
    series = series[series[produit] > 0]

    if series.empty:
        return go.Figure()

    color = COLOR_GREEN if series[produit].iloc[-1] >= series[produit].iloc[0] else COLOR_RED

    fig = go.Figure()
    fig.add_trace(go.Scatter(
        x=series["date"],
        y=series[produit],
        mode="lines+markers",
        line=dict(color=color, width=2.5),
        marker=dict(size=6),
        hovertemplate="<b>%{x|%d %b %Y}</b><br>Valeur : %{y:,.2f}€<extra></extra>",
    ))

    # Ligne de référence premier point
    first_val = float(series[produit].iloc[0])
    fig.add_hline(y=first_val, line_dash="dot",
                  line_color="rgba(128,128,128,0.4)",
                  annotation_text=f"Initial value: {first_val:,.0f}€")

    y_min = float(series[produit].min()) * 0.92
    y_max = float(series[produit].max()) * 1.05
    fig.update_layout(
        title=dict(text=f"Position Value — {produit}", font=dict(size=14), x=0.02),
        xaxis=dict(showgrid=True, gridcolor=COLOR_GRID),
        yaxis=dict(title="Valeur (€)", showgrid=True, gridcolor=COLOR_GRID,
                   ticksuffix="€", range=[y_min, y_max]),
        plot_bgcolor=COLOR_BG, paper_bgcolor=COLOR_BG,
        margin=dict(l=60, r=20, t=50, b=40),
        height=320, showlegend=False,
    )
    return fig


def build_projection(
    total_val: float,
    versement_mensuel: float,
    horizons: List[int] = [10, 20, 30],
    current_invested: Optional[float] = None,
) -> go.Figure:
    """
    Projection DCA sur 10/20/30 ans.
    3 scénarios : pessimiste 5%, base 8%, optimiste 11%
    """
    scenarios = [
        ("Conservative (5%/yr)",  0.05, COLOR_RED),
        ("Base case (8%/yr)",        0.08, COLOR_AMBER),
        ("Optimistic (11%/yr)", 0.11, COLOR_GREEN),
    ]

    fig = go.Figure()
    max_horizon = max(horizons)

    for label, taux_annuel, color in scenarios:
        taux_mensuel = (1 + taux_annuel) ** (1/12) - 1
        valeurs = []
        mois_list = list(range(0, max_horizon * 12 + 1))

        val = total_val
        for m in mois_list:
            if m > 0:
                val = val * (1 + taux_mensuel) + versement_mensuel
            valeurs.append(val)

        annees = [m / 12 for m in mois_list]

        fig.add_trace(go.Scatter(
            x=annees, y=valeurs,
            name=label,
            line=dict(color=color, width=2.5),
            hovertemplate=f"<b>{label}</b><br>Dans %{{x:.0f}} ans : %{{y:,.0f}}€<extra></extra>",
        ))

    # Historical invested capital plus future contributions, with zero future return.
    invested_baseline = float(current_invested) if current_invested is not None else float(total_val)
    val_investi = [invested_baseline + versement_mensuel * m for m in mois_list]
    fig.add_trace(go.Scatter(
        x=annees, y=val_investi,
        name="Current invested + future contributions (0% return)",
        line=dict(color="rgba(128,128,128,0.6)", width=1.5, dash="dot"),
        hovertemplate="Capital investi : %{y:,.0f}€<extra></extra>",
    ))

    # Marqueurs aux horizons clés
    for horizon in horizons:
        fig.add_vline(
            x=horizon,
            line_dash="dash",
            line_color="rgba(128,128,128,0.3)",
            annotation_text=f"{horizon}y",
            annotation_position="top",
        )

    fig.update_layout(
        title=dict(text=f"Projection DCA — versement {versement_mensuel:,.0f}€/mois", font=dict(size=15), x=0.02),
        xaxis=dict(title="Years", showgrid=True, gridcolor=COLOR_GRID),
        yaxis=dict(title="Estimated value (€)", showgrid=True, gridcolor=COLOR_GRID,
                   tickformat=",.0f", ticksuffix="€"),
        plot_bgcolor=COLOR_BG, paper_bgcolor=COLOR_BG,
        margin=dict(l=80, r=20, t=50, b=50),
        hovermode="x unified", height=400,
        legend=dict(orientation="h", x=0, y=1.12),
    )
    return fig


def build_allocation_donut(df: pd.DataFrame) -> go.Figure:
    """Donut allocation actuelle tous produits."""
    last = df.iloc[-1]

    labels, values = [], []

    # ETFs principaux
    col_map = {"iShares S&P 500 Acc": "sp500_val",
               "iShares NASDAQ 100 Acc": "nasdaq_val",
               "HSBC MSCI Emerging Markets Acc": "em_val",
               "VanEck Defense Acc": "defense_val"}
    for nom in ETF_PRINCIPAUX:
        v = float(last.get(col_map[nom], 0) or 0)
        if v > 0:
            labels.append(nom.replace(" Acc", "").replace("iShares ", "").replace("HSBC ", "").replace("VanEck ", ""))
            values.append(v)

    # Autres produits
    for nom in PRODUITS_MAP:
        v = float(last.get(nom, 0) or 0)
        if v > 50:
            short = nom.replace(" Acc", "").replace(" Dist", "").replace("iShares ", "").replace("Xtrackers ", "").replace("Vanguard ", "").replace("VanEck ", "").replace("Global X ", "").replace("L&G ", "").replace("HSBC ", "")
            labels.append(short)
            values.append(v)

    fig = go.Figure(go.Pie(
        labels=labels, values=values,
        hole=0.55,
        textinfo="none",
        hovertemplate="<b>%{label}</b><br>%{value:,.0f}€ (%{percent})<extra></extra>",
        marker=dict(colors=px.colors.qualitative.Set3),
    ))
    fig.update_layout(
        title=dict(text="Allocation actuelle", font=dict(size=15), x=0.02),
        paper_bgcolor=COLOR_BG,
        margin=dict(l=20, r=20, t=50, b=20),
        height=400, showlegend=False,
    )
    return fig


def load_invested_from_dashboard(path: str = EXCEL_PATH) -> Dict:
    """Read canonical invested amounts from the workbook dashboard."""
    invested = {}
    wb = None
    try:
        wb = openpyxl.load_workbook(path, keep_vba=path.lower().endswith('.xlsm'), data_only=True)
        if 'dashboard' not in wb.sheetnames:
            return invested
        ws = wb['dashboard']

        if ws.max_row > MAX_DCA_ROWS or ws.max_column > MAX_DCA_COLUMNS:
            logger.warning(
                "Dashboard dimensions exceed safety limits in %s: %s x %s",
                path,
                ws.max_row,
                ws.max_column,
            )
            return invested

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=3, values_only=True):
            if len(row) < 3:
                continue
            name, inv = row[1], _finite_number(row[2])
            if name and inv is not None and inv > 0:
                raw_name = str(name).strip()
                canonical_name = DASHBOARD_PRODUCT_ALIASES.get(raw_name, raw_name)
                if canonical_name in invested:
                    existing = invested[canonical_name]
                    if existing is None or abs(existing - inv) > 1e-9:
                        logger.warning("Conflicting invested amounts for product %s", canonical_name)
                        invested[canonical_name] = None
                    # Identical duplicate rows are harmless and remain canonical.
                else:
                    invested[canonical_name] = inv
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as exc:
        logger.warning("Could not read dashboard invested amounts from %s: %s", path, exc)
    finally:
        if wb is not None:
            wb.close()
    return invested


def reconcile_latest_snapshot(df: pd.DataFrame) -> Dict:
    """Reconcile the latest mapped position values against workbook total_val.

    A non-zero residual is surfaced instead of silently pretending the allocation
    map is complete (for example after the workbook gains a new product column).
    """
    if df is None or df.empty:
        return {"mapped_value": None, "total_value": None, "residual": None,
                "residual_pct": None, "reconciled": False}
    last = df.iloc[-1]
    mapped = 0.0
    principal_cols = ("sp500_val", "nasdaq_val", "em_val", "defense_val")
    for col in principal_cols:
        value = _finite_number(last.get(col, 0))
        if value is not None:
            mapped += value
    for name in PRODUITS_MAP:
        value = _finite_number(last.get(name, 0))
        if value is not None:
            mapped += value
    total = _finite_number(last.get("total_val"), allow_zero=False)
    if total is None:
        return {"mapped_value": mapped, "total_value": None, "residual": None,
                "residual_pct": None, "reconciled": False}
    residual = total - mapped
    tolerance = max(1.0, total * 0.001)  # €1 or 10 bp of portfolio, whichever is larger
    return {
        "mapped_value": round(mapped, 2),
        "total_value": round(total, 2),
        "residual": round(residual, 2),
        "residual_pct": round(residual / total, 6),
        "reconciled": abs(residual) <= tolerance,
    }


def get_all_produits(df: pd.DataFrame) -> List[str]:
    """Liste tous les produits disponibles, triés par valeur décroissante."""
    last = df.iloc[-1]
    produits_vals = []

    col_map = {
        "iShares S&P 500 Acc":            "sp500_val",
        "iShares NASDAQ 100 Acc":          "nasdaq_val",
        "HSBC MSCI Emerging Markets Acc":  "em_val",
        "VanEck Defense Acc":              "defense_val",
    }
    for nom, col in col_map.items():
        v = float(last.get(col, 0) or 0)
        if v > 0:
            produits_vals.append((nom, v))

    for nom in PRODUITS_MAP:
        v = float(last.get(nom, 0) or 0)
        if v > 0:
            produits_vals.append((nom, v))

    produits_vals.sort(key=lambda x: -x[1])
    return [p[0] for p in produits_vals]


def get_produit_stats(df: pd.DataFrame, produit: str, excel_path: str = EXCEL_PATH) -> Dict:
    """Return product P&L using dashboard invested capital as the source of truth."""
    col_map = {
        "iShares S&P 500 Acc": ("sp500_inv", "sp500_val"),
        "iShares NASDAQ 100 Acc": ("nasdaq_inv", "nasdaq_val"),
        "HSBC MSCI Emerging Markets Acc": ("em_inv", "em_val"),
        "VanEck Defense Acc": ("defense_inv", "defense_val"),
    }
    last = df.iloc[-1]
    value_col = col_map.get(produit, (None, produit))[1]
    valeur = _finite_number(last.get(value_col, 0))
    if valeur is None:
        return {"investi": None, "valeur": None, "pnl": None, "rendement": None,
                "invested_available": False, "error": f"Invalid current value for {produit}."}

    dashboard_inv = load_invested_from_dashboard(excel_path)
    dashboard_value = dashboard_inv.get(produit)
    data_value = _finite_number(last.get(col_map[produit][0])) if produit in col_map else None

    if dashboard_value is not None and data_value is not None and data_value > 0:
        tolerance = max(1.0, 0.01 * dashboard_value)
        if abs(dashboard_value - data_value) > tolerance:
            return {"investi": None, "valeur": valeur, "pnl": None, "rendement": None,
                    "invested_available": False,
                    "error": f"Invested-capital mismatch for {produit}: dashboard and data sheet disagree."}

    investi = dashboard_value if dashboard_value is not None else data_value
    if investi is None or investi <= 0:
        return {"investi": None, "valeur": valeur, "pnl": None, "rendement": None,
                "invested_available": False,
                "error": f"Invested amount unavailable for {produit}; no market-value fallback is allowed."}

    pnl = valeur - investi
    return {"investi": float(investi), "valeur": float(valeur), "pnl": pnl,
            "rendement": pnl / investi, "invested_available": True, "error": None}


def build_ai_portfolio_prompt(stats: Dict, df: pd.DataFrame, dataset_mode: str = "personal") -> str:
    """Build a complete, bounded portfolio prompt; no hidden truncation of concentration data."""
    last = df.iloc[-1]
    positions = []
    col_map = {"iShares S&P 500 Acc": "sp500_val", "iShares NASDAQ 100 Acc": "nasdaq_val",
               "HSBC MSCI Emerging Markets Acc": "em_val", "VanEck Defense Acc": "defense_val"}
    for nom, col in col_map.items():
        v = _finite_number(last.get(col, 0))
        if v is not None and v > 0:
            positions.append((nom, v))
    for nom in PRODUITS_MAP:
        v = _finite_number(last.get(nom, 0))
        if v is not None and v > 0:
            positions.append((nom, v))
    positions.sort(key=lambda x: x[1], reverse=True)
    total_positions = sum(v for _, v in positions)
    reconciliation = reconcile_latest_snapshot(df)
    lines = []
    for nom, value in positions:
        weight = value / stats["total_val"] if stats.get("total_val", 0) > 0 else 0
        lines.append(f"  - {nom}: {value:,.0f}€ ({weight*100:.1f}% of total portfolio)")
    positions_str = "\n".join(lines) if lines else "  - No position breakdown available"
    unallocated = reconciliation.get("residual")
    if unallocated is None:
        unallocated = stats.get("total_val", 0) - total_positions
    xirr_str = f"{stats['xirr']*100:+.1f}%" if stats.get("xirr") is not None else "N/A"

    is_demo = str(dataset_mode).lower() == "demo"
    allocation_reconciled = bool(reconciliation.get("reconciled", False))
    if is_demo:
        heading = "ANONYMISED DEMO WORKBOOK DATA"
        dataset_note = (
            "This is the repository demo dataset, not the user's personal portfolio. "
            "Analyse it only as an example dataset."
        )
    else:
        heading = "PERSONAL WORKBOOK DATA"
        dataset_note = "This is the locally loaded personal workbook dataset."

    allocation_status = (
        "RECONCILED"
        if allocation_reconciled
        else "INCOMPLETE / UNRECONCILED"
    )

    return f"""You are a portfolio analyst reviewing a long-term DCA portfolio dataset.

═══ {heading} ═══
Dataset mode      : {'DEMO' if is_demo else 'PERSONAL'}
Dataset note      : {dataset_note}
Total invested    : {stats['total_inv']:,.0f}€
Current value     : {stats['total_val']:,.0f}€
Total P&L         : {stats['pnl']:+,.0f}€ ({stats['pnl_pct']*100:+.1f}%)
Money-weighted return (XIRR): {xirr_str}
Observed period   : {stats['date_debut']} → {stats['date_fin']}
Portfolio snapshots: {stats['n_observations']}
Contribution events inferred from cumulative invested capital: {stats.get('contribution_events', 'N/A')}

MAPPED POSITIONS FROM THE LATEST SNAPSHOT:
{positions_str}
Allocation mapping status: {allocation_status}
Residual/unmapped value versus portfolio total: {unallocated:,.0f}€

═══ DATA INTEGRITY RULES ═══
- Total invested, total value, P&L, XIRR and dates above are workbook-derived figures.
- Treat the mapped position list as complete only when Allocation mapping status is RECONCILED.
- If allocation mapping is INCOMPLETE / UNRECONCILED, do not infer the identity, sector or risk of the residual/unmapped value.
- No live prices, current macro data, news, benchmark returns, risk-free rate, correlations or forward-looking forecasts are supplied.
- Do not invent current market conditions, sector outlooks, valuation multiples or catalysts.
- You may discuss concentration/diversification from the supplied weights, but label forward-looking return/risk statements as scenarios or qualitative hypotheses.
- If a requested fact requires data not supplied here, state that it is unavailable.

Be direct and quantitative. Separate workbook observations from assumptions."""

