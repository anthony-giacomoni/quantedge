from datetime import datetime
from pathlib import Path

import pandas as pd

from modules.dca import calculate_xirr, compute_global_stats, load_excel, get_produit_stats, DEMO_EXCEL_PATH


def test_xirr_one_year_single_contribution():
    df = pd.DataFrame({
        "date": [pd.Timestamp("2025-01-01"), pd.Timestamp("2026-01-01")],
        "total_inv": [100.0, 100.0],
        "total_val": [100.0, 110.0],
    })
    xirr = calculate_xirr(df)
    assert xirr is not None
    assert abs(xirr - 0.10) < 0.001


def test_demo_dataset_is_reproducible_and_has_xirr():
    assert Path(DEMO_EXCEL_PATH).exists()
    df = load_excel(DEMO_EXCEL_PATH)
    assert df is not None and not df.empty
    stats = compute_global_stats(df)
    assert stats["total_inv"] > 0
    assert stats["xirr"] is not None


def test_dca_source_prefers_private_workbook(monkeypatch, tmp_path):
    import modules.dca as dca

    personal = tmp_path / "simu_invest.xlsm"
    demo = tmp_path / "demo.xlsx"
    personal.write_bytes(b"personal")
    demo.write_bytes(b"demo")

    monkeypatch.setattr(dca, "EXCEL_PATH", str(personal))
    monkeypatch.setattr(dca, "DEMO_EXCEL_PATH", str(demo))
    path, mode = dca.resolve_dca_source()
    assert path == str(personal)
    assert mode == "personal"


def test_dca_source_uses_demo_only_when_private_missing(monkeypatch, tmp_path):
    import modules.dca as dca

    personal = tmp_path / "missing.xlsm"
    demo = tmp_path / "demo.xlsx"
    demo.write_bytes(b"demo")

    monkeypatch.setattr(dca, "EXCEL_PATH", str(personal))
    monkeypatch.setattr(dca, "DEMO_EXCEL_PATH", str(demo))
    path, mode = dca.resolve_dca_source()
    assert path == str(demo)
    assert mode == "demo"


def _write_dashboard_workbook(path, rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "dashboard"
    for excel_row, name, invested in rows:
        ws.cell(excel_row, 2, name)
        ws.cell(excel_row, 3, invested)
    wb.save(path)


def test_gold_dashboard_alias_reconciles_invested_amount(tmp_path):
    path = tmp_path / "portfolio.xlsx"
    _write_dashboard_workbook(path, [(20, "iShares Physical Gold Acc (environ)", 1586.0)])
    df = pd.DataFrame({"iShares Physical Gold Acc": [285.66, 2090.36]})
    stats = get_produit_stats(df, "iShares Physical Gold Acc", excel_path=str(path))
    assert stats["investi"] == 1586.0
    assert abs(stats["pnl"] - 504.36) < 1e-9
    assert abs(stats["rendement"] - (504.36 / 1586.0)) < 1e-9
    assert stats["invested_available"] is True


def test_missing_product_invested_never_falls_back_to_market_value(tmp_path):
    path = tmp_path / "portfolio.xlsx"
    _write_dashboard_workbook(path, [])
    df = pd.DataFrame({"Unknown Product": [100.0, 700.0]})
    stats = get_produit_stats(df, "Unknown Product", excel_path=str(path))
    assert stats["investi"] is None
    assert stats["pnl"] is None
    assert stats["rendement"] is None
    assert stats["invested_available"] is False


def test_dca_rows_are_reported_as_snapshots_not_contributions():
    df = pd.DataFrame({
        "date": [pd.Timestamp("2025-01-01"), pd.Timestamp("2025-02-01"), pd.Timestamp("2025-03-01")],
        "total_inv": [1000.0, 2000.0, 3000.0],
        "total_val": [1000.0, 2050.0, 3150.0],
    })
    stats = compute_global_stats(df)
    assert stats["n_observations"] == 3
    assert "n_versements" not in stats
    assert stats["avg_monthly_contribution"] > 0


def _make_invalid_workbook(path, total_inv=100.0, total_val=110.0):
    import openpyxl
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="data"
    ws.append(["date"]+[None]*11)
    row=[__import__('datetime').datetime(2026,1,1), 10,10,10,10,total_inv,0,20,20,20,20,total_val]
    ws.append(row); wb.save(path); wb.close()


def test_dca_rejects_nan_and_negative_numeric_inputs(tmp_path):
    from modules.dca import load_excel, DCADataError
    import pytest
    p=tmp_path/"nan.xlsx"; _make_invalid_workbook(p,total_inv=float('nan'))
    with pytest.raises(DCADataError): load_excel(str(p))
    p2=tmp_path/"negative.xlsx"; _make_invalid_workbook(p2,total_inv=-1)
    with pytest.raises(DCADataError): load_excel(str(p2))


def test_dca_rejects_decreasing_cumulative_invested(tmp_path):
    import openpyxl, pytest
    from modules.dca import load_excel, DCADataError
    p=tmp_path/"decrease.xlsx"; wb=openpyxl.Workbook(); ws=wb.active; ws.title="data"
    ws.append(["date"]+[None]*11)
    for dt,inv,val in [(__import__('datetime').datetime(2026,1,1),100,110),(__import__('datetime').datetime(2026,2,1),90,105)]:
        ws.append([dt,10,10,10,10,inv,0,20,20,20,20,val])
    wb.save(p); wb.close()
    with pytest.raises(DCADataError): load_excel(str(p))


def test_product_stats_rejects_dashboard_data_sheet_mismatch(monkeypatch):
    import pandas as pd
    from modules import dca
    df=pd.DataFrame([{"sp500_inv":100.0,"sp500_val":120.0}])
    monkeypatch.setattr(dca,"load_invested_from_dashboard",lambda path:{"iShares S&P 500 Acc":150.0})
    stats=dca.get_produit_stats(df,"iShares S&P 500 Acc","fake.xlsx")
    assert stats["invested_available"] is False
    assert "mismatch" in stats["error"]


def test_portfolio_ai_prompt_includes_all_positions_and_guardrail():
    import pandas as pd
    from modules import dca
    row={"sp500_val":100,"nasdaq_val":100,"em_val":100,"defense_val":100}
    for i,name in enumerate(dca.PRODUITS_MAP): row[name]=100+i
    df=pd.DataFrame([row])
    stats={"total_inv":1000,"total_val":3000,"pnl":2000,"pnl_pct":2.0,"xirr":0.1,"date_debut":"01/01/2025","date_fin":"01/01/2026","n_observations":10,"contribution_events":5}
    prompt=dca.build_ai_portfolio_prompt(stats,df)
    assert prompt.count("  - ") >= len(dca.PRODUITS_MAP)+4
    assert "No live prices, current macro data, news" in prompt
    assert "Residual/unmapped value" in prompt


def test_latest_snapshot_reconciliation_detects_unmapped_value():
    from modules import dca
    import pandas as pd
    row = {
        "total_val": 1000.0,
        "sp500_val": 400.0, "nasdaq_val": 300.0, "em_val": 100.0, "defense_val": 100.0,
    }
    for name in dca.PRODUITS_MAP:
        row[name] = 0.0
    df = pd.DataFrame([row])
    rec = dca.reconcile_latest_snapshot(df)
    assert rec["mapped_value"] == 900.0
    assert rec["residual"] == 100.0
    assert rec["reconciled"] is False


def test_latest_snapshot_reconciliation_accepts_exact_mapping():
    from modules import dca
    import pandas as pd
    row = {
        "total_val": 1000.0,
        "sp500_val": 400.0, "nasdaq_val": 300.0, "em_val": 100.0, "defense_val": 200.0,
    }
    for name in dca.PRODUITS_MAP:
        row[name] = 0.0
    df = pd.DataFrame([row])
    rec = dca.reconcile_latest_snapshot(df)
    assert rec["residual"] == 0.0
    assert rec["reconciled"] is True


def test_dashboard_conflicting_aliases_remain_unavailable_after_third_duplicate(tmp_path):
    import openpyxl
    from modules.dca import load_invested_from_dashboard
    path = tmp_path / "conflict.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "dashboard"
    ws.append([None, "iShares Physical Gold Acc", 1000.0])
    ws.append([None, "iShares Physical Gold Acc (environ)", 1200.0])
    ws.append([None, "iShares Physical Gold Acc", 1300.0])
    wb.save(path)
    wb.close()
    invested = load_invested_from_dashboard(str(path))
    assert invested["iShares Physical Gold Acc"] is None

def test_portfolio_ai_prompt_does_not_claim_reconciliation_when_mapping_has_residual():
    import pandas as pd
    from modules import dca

    row = {
        "total_val": 1000.0,
        "sp500_val": 400.0,
        "nasdaq_val": 300.0,
        "em_val": 100.0,
        "defense_val": 100.0,
    }

    for name in dca.PRODUITS_MAP:
        row[name] = 0.0

    df = pd.DataFrame([row])

    stats = {
        "total_inv": 900.0,
        "total_val": 1000.0,
        "pnl": 100.0,
        "pnl_pct": 100.0 / 900.0,
        "xirr": 0.05,
        "date_debut": "01/01/2025",
        "date_fin": "01/01/2026",
        "n_observations": 2,
        "contribution_events": 1,
    }

    prompt = dca.build_ai_portfolio_prompt(stats, df)

    assert (
        "Allocation mapping status: INCOMPLETE / UNRECONCILED"
        in prompt
    )
    assert "MAPPED POSITIONS FROM THE LATEST SNAPSHOT" in prompt
    assert "ALL POSITIONS RECORDED IN THE LATEST SNAPSHOT" not in prompt
    assert "RECONCILED PERSONAL WORKBOOK DATA" not in prompt
    assert "do not infer the identity, sector or risk" in prompt

def test_ai_prompt_reconciled_population_includes_small_mapped_holdings():
    import pandas as pd
    from modules import dca

    row = {
        "sp500_val": 1000.0,
        "nasdaq_val": 0.0,
        "em_val": 0.0,
        "defense_val": 0.0,
    }

    for name in dca.PRODUITS_MAP:
        row[name] = 50.0

    total = 1000.0 + 50.0 * len(dca.PRODUITS_MAP)
    row["total_val"] = total

    df = pd.DataFrame([row])

    stats = {
        "total_inv": total,
        "total_val": total,
        "pnl": 0.0,
        "pnl_pct": 0.0,
        "xirr": None,
        "date_debut": "01/01/2026",
        "date_fin": "01/01/2026",
        "n_observations": 1,
        "contribution_events": 1,
    }

    reconciliation = dca.reconcile_latest_snapshot(df)

    assert reconciliation["reconciled"] is True
    assert reconciliation["residual"] == 0.0

    prompt = dca.build_ai_portfolio_prompt(
        stats,
        df,
    )

    assert "Allocation mapping status: RECONCILED" in prompt

    # Every positive mapped holding participating in reconciliation must also
    # be represented in the Claude allocation payload.
    for name in dca.PRODUITS_MAP:
        assert name in prompt

