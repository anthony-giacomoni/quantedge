from pathlib import Path

import pandas as pd
from openpyxl import Workbook


def test_history_provider_duplicate_session_is_deduplicated():
    from utils.market_data import _sanitize_history_df

    dates = pd.to_datetime(
        [
            "2026-09-01",
            "2026-09-02",
            "2026-09-02",
            "2026-09-03",
        ]
    )

    df = pd.DataFrame(
        {
            "Open": [10.0, 11.0, 11.5, 12.0],
            "High": [11.0, 12.0, 12.0, 13.0],
            "Low": [9.0, 10.0, 11.0, 11.0],
            "Close": [10.5, 11.5, 11.7, 12.5],
            "Volume": [100, 110, 120, 130],
        },
        index=dates,
    )

    out = _sanitize_history_df("AAPL", df)

    assert out is not None
    assert out.index.is_unique
    assert len(out) == 3

    # Last provider record wins for the duplicated session.
    assert (
        float(out.loc[pd.Timestamp("2026-09-02"), "Close"])
        == 11.7
    )


def test_validator_detects_inline_generic_secret(tmp_path):
    from scripts.validate_public_repo import validate

    folder = tmp_path / "modules"
    folder.mkdir()

    # Construct dynamically so this public test source does not itself
    # contain a secret-looking literal.
    secret = (
        "OPENAI_"
        + "API_KEY="
        + "REAL_PRIVATE_VALUE_123456"
    )

    (folder / "private.py").write_text(
        f'cfg = "{secret}"\n'
    )

    problems = validate(tmp_path)

    assert any(
        "OPENAI_API_KEY" in problem
        for problem in problems
    )


def test_dca_rejects_sparse_sheet_with_huge_declared_row(tmp_path):
    from modules.dca import load_excel, DCADataError

    path = tmp_path / "sparse.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "data"

    ws.append(
        [f"c{i}" for i in range(1, 69)]
    )

    # Produces a tiny archive but an enormous worksheet dimension.
    ws["A10001"] = "2026-01-01"

    wb.save(path)
    wb.close()

    try:
        load_excel(str(path))
    except DCADataError as exc:
        assert "dimensions exceed safety limits" in str(exc)
    else:
        raise AssertionError(
            "Sparse oversized worksheet was not rejected"
        )


def test_dashboard_sparse_dimension_fails_closed(tmp_path):
    from modules.dca import load_invested_from_dashboard

    path = tmp_path / "dashboard_sparse.xlsx"

    wb = Workbook()

    ws = wb.active
    ws.title = "data"
    ws.append(["date"])

    dash = wb.create_sheet("dashboard")
    dash["A10001"] = "sentinel"

    wb.save(path)
    wb.close()

    assert load_invested_from_dashboard(str(path)) == {}
