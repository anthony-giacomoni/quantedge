from datetime import datetime

import pandas as pd
from openpyxl import Workbook


def test_portfolio_rejects_stale_quote_for_open_pnl(
    monkeypatch,
    tmp_path,
):
    from utils.db import init_database, insert_trade
    from modules import portfolio

    db = str(tmp_path / "q.db")
    init_database(db)

    insert_trade(
        {
            "ticker": "AAPL",
            "direction": "LONG",
            "qty": 10,
            "entry_price": 100,
            "entry_date": "2026-01-01",
            "exit_price": None,
            "exit_date": None,
            "pnl": None,
            "invested": 1000,
            "status": "OPEN",
            "sector": "Technology",
            "notes": "",
            "instrument_type": "DIRECT",
            "marking_ticker": None,
        },
        db,
    )

    monkeypatch.setattr(
        portfolio,
        "get_quote",
        lambda ticker: {
            "price": 120,
            "currency": "USD",
            "provider_reported_currency": None,
            "stale": True,
            "error": None,
            "source": "stale test quote",
            "timestamp": "2026-01-01",
        },
    )

    monkeypatch.setattr(
        portfolio,
        "get_fx_quote_to_eur",
        lambda currency: {
            "rate": 1.0,
            "stale": False,
            "error": None,
            "source": "identity",
            "timestamp": None,
        },
    )

    row = portfolio.load_enriched_trades(db).iloc[0]

    assert pd.isna(row["pnl"])
    assert bool(row["pricing_stale"]) is True


def test_compute_open_pnl_rejects_stale_quote(monkeypatch):
    from utils import market_data

    monkeypatch.setattr(
        market_data,
        "get_quote",
        lambda ticker: {
            "price": 120,
            "currency": "USD",
            "stale": True,
            "error": None,
        },
    )

    monkeypatch.setattr(
        market_data,
        "get_fx_quote_to_eur",
        lambda currency: {
            "rate": 1.0,
            "stale": False,
            "error": None,
        },
    )

    assert (
        market_data.compute_open_pnl(
            "AAPL",
            100,
            10,
            "LONG",
        )
        is None
    )


def test_ai_snapshot_hides_stale_quote_price(monkeypatch):
    from modules import ai_analysis

    monkeypatch.setattr(
        ai_analysis,
        "get_quote",
        lambda ticker: {
            "price": 123.0,
            "stale": True,
            "error": None,
            "source": "old quote",
            "currency": "USD",
            "timestamp": "old",
        },
    )

    monkeypatch.setattr(
        ai_analysis,
        "get_price_history_with_meta",
        lambda *args, **kwargs: {
            "data": None,
            "source": None,
            "as_of": None,
            "error": "stale",
        },
    )

    monkeypatch.setattr(
        ai_analysis,
        "get_fundamentals",
        lambda ticker: {
            "_available": False,
            "_source": "yfinance",
        },
    )

    out = ai_analysis.collect_market_snapshot("AAPL")

    assert out["current_price"] is None
    assert out["quote_stale"] is True


def test_nonempty_unusable_macro_payload_falls_back(
    monkeypatch,
):
    from utils import event_data

    monkeypatch.setattr(
        event_data,
        "_eodhd_macro_cached",
        lambda *args, **kwargs: (
            {
                "date": "1900-01-01",
                "country": "US",
                "type": "Old event",
            },
            None,
        ),
    )

    monkeypatch.setattr(
        event_data,
        "_official_macro_events",
        lambda *args, **kwargs: (
            [
                {
                    "date": "2026-09-10",
                    "country": "US",
                    "event": "Official fallback",
                    "impact": "Major",
                    "focus": "Macro",
                    "estimate": None,
                    "previous": None,
                    "actual": None,
                    "unit": None,
                    "source": "Official",
                }
            ],
            [],
        ),
    )

    out = event_data.get_upcoming_macro(
        30,
        ["US"],
    )

    assert out["source"].startswith(
        "Official macro calendars"
    )


def test_nonempty_unusable_earnings_payload_falls_back(
    monkeypatch,
):
    from utils import event_data

    monkeypatch.setattr(
        event_data,
        "_eodhd_earnings_cached",
        lambda *args, **kwargs: (
            [
                {
                    "report_date": "1900-01-01",
                    "code": "AAPL.US",
                }
            ],
            None,
        ),
    )

    monkeypatch.setattr(
        event_data,
        "_yfinance_earnings",
        lambda *args, **kwargs: (
            [
                {
                    "date": "2026-09-15",
                    "ticker": "AAPL",
                    "before_after_market": None,
                    "eps_estimate": None,
                    "currency": None,
                    "source": "yfinance Earnings Calendar",
                }
            ],
            [],
        ),
    )

    out = event_data.get_upcoming_earnings(
        ["AAPL"],
        30,
    )

    assert out["source"] == "yfinance Earnings Calendar"


def test_demo_validator_scans_workbook_cells_for_private_content(
    tmp_path,
):
    from scripts.validate_public_repo import validate

    demo = (
        tmp_path
        / "examples"
        / "simu_invest_demo.xlsx"
    )

    demo.parent.mkdir(parents=True)

    wb = Workbook()
    ws = wb.active

    # Concatenated deliberately so the source test itself
    # does not contain a real-looking public secret/PII literal.
    ws["A1"] = (
        "private.person"
        + "@"
        + "private-domain.test"
    )

    ws["A2"] = (
        "EODHD_"
        + "API_KEY="
        + "REAL_PRIVATE_VALUE_12345"
    )

    wb.save(demo)
    wb.close()

    problems = validate(tmp_path)

    assert any(
        "personal email" in problem
        for problem in problems
    )

    assert any(
        "EODHD_API_KEY" in problem
        for problem in problems
    )


def test_missing_product_cell_is_unknown_not_zero(
    tmp_path,
):
    from modules.dca import load_excel

    path = tmp_path / "blank.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "data"

    ws.append(
        [f"c{i}" for i in range(1, 69)]
    )

    row = [0] * 68

    row[0] = datetime(2025, 1, 1)
    row[5] = 1000.0
    row[11] = 2000.0

    # Physical Gold market-value cell intentionally missing.
    row[37] = None

    ws.append(row)

    wb.save(path)
    wb.close()

    df = load_excel(str(path))

    assert pd.isna(
        df.iloc[-1]["iShares Physical Gold Acc"]
    )
